"""Build reproducible DOCX -> DOCX fixtures and run the actual Celsius pipeline.

Run from the project root: python scripts/document_fill_demo.py --output-dir <folder>
Fixtures are fictional and the expected document is authored independently.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.shared import Cm, Pt


def create_demo(root: Path, *, fixtures_only: bool = False) -> dict:
    root.mkdir(parents=True, exist_ok=True)
    if any(root.iterdir()):
        raise FileExistsError("Use uma pasta vazia para nao sobrescrever os exemplos.")
    source = Document()
    source.add_heading("Informações do aluno", 0)
    facts = [
        ("Nome do aluno", "João da Silva"),
        ("Data de nascimento", "10/04/2017"),
        ("Escola", "Escola Municipal Exemplo"),
        ("Necessita adaptação", "Sim"),
        ("Observações", "Participa das atividades com apoio visual."),
    ]
    for label, value in facts:
        source.add_paragraph(f"{label}: {value}")
    source.save(root / "origem.docx")

    template = Document()
    normal = template.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    section = template.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.header.paragraphs[0].text = "Secretaria Municipal de Educação"
    section.footer.paragraphs[0].text = "Modelo padronizado de acompanhamento"
    template.add_heading("Relatório de acompanhamento", 0)
    template.add_paragraph("Identificação e observações para acompanhamento pedagógico.")
    table = template.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    labels = ["Nome:", "Data de nascimento:", "Unidade escolar:",
              "Necessita adaptação:", "Observações:", "Telefone:"]
    for index, label in enumerate(labels):
        table.cell(index, 0).paragraphs[0].add_run(label).bold = True
        table.cell(index, 1).paragraphs[0].add_run("").italic = True
    table.cell(3, 1).paragraphs[0].runs[0].text = "( ) Sim ( ) Não"
    template.save(root / "modelo.docx")

    # Expected result does not call the implementation under test.
    expected = Document(root / "modelo.docx")
    expected_values = ["João da Silva", "10/04/2017", "Escola Municipal Exemplo",
                       "(X) Sim ( ) Não", "Participa das atividades com apoio visual.", ""]
    for index, value in enumerate(expected_values):
        if index == 3:
            expected.tables[0].cell(index, 1).paragraphs[0].runs[0].text = value
        elif value:
            expected.tables[0].cell(index, 1).paragraphs[0].add_run(value).italic = True
    expected.save(root / "resultado_esperado.docx")
    if fixtures_only:
        return {"written": False, "output": "", "applied_count": 0, "still_open": []}
    return validate_demo(root)


def validate_demo(root: Path) -> dict:
    from core.document_pipeline import fill_from_sources

    result = fill_from_sources(root / "modelo.docx", [root / "origem.docx"],
                               root / "resultado.docx", use_llm=False)
    actual = Document(result["output"])
    expected = Document(root / "resultado_esperado.docx")
    expected_values = [expected.tables[0].cell(i, 1).text for i in range(6)]
    assert [actual.tables[0].cell(i, 1).text for i in range(6)] == expected_values
    with ZipFile(root / "modelo.docx") as original, ZipFile(root / "resultado.docx") as filled:
        for name in original.namelist():
            if name != "word/document.xml":
                assert original.read(name) == filled.read(name), name
    assert "Telefone:" in result["still_open"]
    (root / "validacao.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--fixtures-only", action="store_true")
    parser.add_argument("--validate-only", action="store_true")
    arguments = parser.parse_args()
    result = (validate_demo(arguments.output_dir.resolve()) if arguments.validate_only
              else create_demo(arguments.output_dir.resolve(), fixtures_only=arguments.fixtures_only))
    print(json.dumps({"written": result["written"], "output": result["output"],
                      "applied_count": result["applied_count"], "still_open": result["still_open"]},
                     ensure_ascii=True))
