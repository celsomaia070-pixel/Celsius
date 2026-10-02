from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path

from processors.base import ProcessadorArquivo, truncar_texto

_MAX_ROWS_PER_SHEET = 10_000
_MAX_COLUMNS_PER_ROW = 60


class ProcessadorXLSX(ProcessadorArquivo):
    extensoes_suportadas = [".xlsx"]

    @classmethod
    def processar(cls, caminho: str, base_dir: Path | None = None) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return (
                "Formato .xlsx nao disponivel: instale openpyxl (pip install celsius[documents])."
            )

        path = cls._validar_caminho(caminho, base_dir)
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            partes: list[str] = []
            for sheet in workbook.worksheets:
                partes.append(cls._processar_sheet(sheet))
            texto = "\n".join(part for part in partes if part.strip()).strip()
        finally:
            workbook.close()

        texto = texto.strip()
        if not texto:
            return "Planilha vazia."
        return truncar_texto(texto)

    @classmethod
    def _processar_sheet(cls, sheet) -> str:
        sheet_name = getattr(sheet, "title", "Planilha")
        linhas: list[str] = []
        for idx_linha, valores_linha in enumerate(sheet.iter_rows(values_only=True), start=1):
            if idx_linha > _MAX_ROWS_PER_SHEET:
                linhas.append("... [linhas restantes omitidas] ...")
                break
            celulas = [
                cls._formatar_celula(valor) for valor in valores_linha[:_MAX_COLUMNS_PER_ROW]
            ]
            if any(c for c in celulas):
                linhas.append(" | ".join(celulas))
        if not linhas:
            return ""
        return f"--- Planilha: {sheet_name} ---\n" + "\n".join(linhas) + "\n--- Fim ---"

    @staticmethod
    def _formatar_celula(valor) -> str:
        if valor is None:
            return ""
        if isinstance(valor, datetime):
            return valor.isoformat(sep=" ", timespec="minutes")
        if isinstance(valor, (date, time)):
            return str(valor)
        if isinstance(valor, float):
            if valor.is_integer():
                return str(int(valor))
            return f"{valor:.4f}".rstrip("0").rstrip(".")
        return str(valor)
