import pytest
import time
from docx import Document
from fastapi.testclient import TestClient
from pathlib import Path

from core.chat_attachments import AttachmentError
from core.chat_outputs import OutputAttachmentStore, collect_outputs, register_output
from core.chat_service import ChatCoordinator
from core.conversations import ConversationManager
from core.modules import MODULE_CHAT, MODULE_INVENTORY, MODULE_SETTINGS
from core.settings import Settings
from core.web_api import EventHub, create_app

ALLOWED = {".docx", ".txt", ".json"}


@pytest.fixture
def web_settings(tmp_path):
    settings = Settings(
        base_dir=tmp_path,
        data_dir=tmp_path / "data",
        resources_dir=tmp_path / "resources",
        logs_dir=tmp_path / "logs",
    )
    settings.mobile.pairing_token = "test-pairing-token"
    settings.customer.company_name = "Empresa Teste"
    settings.customer.company_sector = "Comercio"
    settings.modules.set_enabled([MODULE_CHAT, MODULE_INVENTORY, MODULE_SETTINGS])
    return settings


def _headers() -> dict[str, str]:
    return {"Authorization": "Bearer test-pairing-token"}


def _store(tmp_path: Path, **kwargs) -> OutputAttachmentStore:
    return OutputAttachmentStore(
        root=tmp_path / "chat_outputs",
        allowed_extensions=ALLOWED,
        max_bytes=8 * 1024 * 1024,
        **kwargs,
    )


def _docx(path: Path, text: str = "conteudo") -> Path:
    document = Document()
    document.add_paragraph(text)
    document.save(path)
    return path


# --- store -----------------------------------------------------------------


def test_generated_file_survives_after_the_turn(tmp_path: Path):
    store = _store(tmp_path)
    generated = _docx(tmp_path / "peenchido.docx")

    with collect_outputs(store) as produced:
        stored = register_output(generated)

    assert stored is not None
    assert produced == [stored.public_dict()]
    # The source file is untouched and the copy is addressable by id.
    assert generated.is_file()
    assert store.get(stored.id).path.is_file()
    assert store.get(stored.id).name == "peenchido.docx"


def test_registration_outside_a_chat_turn_is_a_noop(tmp_path: Path):
    generated = _docx(tmp_path / "soltinho.docx")

    assert register_output(generated) is None


def test_download_survives_a_restart(tmp_path: Path):
    store = _store(tmp_path)
    with collect_outputs(store) as produced:
        stored = register_output(_docx(tmp_path / "plano.docx"))

    # A fresh store has no in-memory map, exactly like after a server restart.
    reopened = _store(tmp_path)

    found = reopened.get(stored.id)
    assert found.name == "plano.docx"
    assert produced[0]["id"] == found.id


def test_store_rejects_unknown_and_malformed_ids(tmp_path: Path):
    store = _store(tmp_path)
    with collect_outputs(store) as produced:
        stored = register_output(_docx(tmp_path / "plano.docx"))

    for bad in ("", "../../etc/passwd", "ZZZZ", stored.id[:-1] + "z", "a" * 40):
        with pytest.raises(AttachmentError):
            store.get(bad)


def test_store_refuses_unsupported_format_and_missing_file(tmp_path: Path):
    store = _store(tmp_path)

    with pytest.raises(AttachmentError):
        store.register(_docx(tmp_path / "malware.exe"), "malware.exe")
    with pytest.raises(AttachmentError):
        store.register(tmp_path / "nao_existe.docx")


def test_store_enforces_the_size_limit(tmp_path: Path):
    store = OutputAttachmentStore(
        root=tmp_path / "chat_outputs",
        allowed_extensions=ALLOWED,
        max_bytes=8,
    )

    with pytest.raises(AttachmentError):
        store.register(_docx(tmp_path / "grande.docx"))


def test_prune_removes_files_older_than_the_retention_window(tmp_path: Path):
    store = OutputAttachmentStore(
        root=tmp_path / "chat_outputs",
        allowed_extensions=ALLOWED,
        max_bytes=1024 * 1024,
        max_age_days=30,
    )
    with collect_outputs(store) as produced:
        stored = register_output(_docx(tmp_path / "antigo.docx"))
    stale = store.get(stored.id).path
    import os
    import time

    ancient = time.time() - (60 * 86_400)
    os.utime(stale, (ancient, ancient))

    removed = store.prune()

    assert removed == 1
    assert not stale.exists()


def _wait_for_job(client: TestClient, job_id: str) -> dict:
    for _attempt in range(200):
        job = client.get(f"/api/v1/chat/jobs/{job_id}", headers=_headers()).json()["job"]
        if job["status"] in {"completed", "failed", "cancelled"}:
            return job
        time.sleep(0.02)
    raise AssertionError("a tarefa de chat nao terminou")


# --- coordinator -----------------------------------------------------------


def test_assistant_reply_persists_and_publishes_generated_file(web_settings, tmp_path):
    web_settings.features.memory = False
    generated = _docx(tmp_path / "plano_preenchido.docx", "Arthur Medeiros Gomes")

    def responder(_prompt, *, fn_chunk, **_kwargs):
        fn_chunk("Preenchido.")
        from core.chat_outputs import register_output

        register_output(generated)
        return "Preenchido."

    hub = EventHub()
    coordinator = ChatCoordinator(
        settings=web_settings,
        event_hub=hub,
        conversation_manager=ConversationManager(tmp_path / "conversations"),
        responder=responder,
        ensure_model_ready=lambda _status: None,
    )
    app = create_app(settings=web_settings, event_hub=hub, chat_coordinator=coordinator)

    with TestClient(app) as client:
        accepted = client.post(
            "/api/v1/chat/messages",
            headers=_headers(),
            json={"message": "preencha o documento"},
        )
        job_id = accepted.json()["job"]["id"]
        job = _wait_for_job(client, job_id)
        assert job["status"] == "completed"

        attachments = job["attachments"]
        assert len(attachments) == 1
        assert attachments[0]["name"] == "plano_preenchido.docx"

        conversation = client.get(
            f"/api/v1/chat/conversations/{job['conversation_id']}", headers=_headers()
        ).json()["conversation"]
        reply = [m for m in conversation["messages"] if m["role"] == "assistant"][-1]
        assert reply["metadata"]["attachments"] == attachments

        # And the file is really downloadable through the API.
        download = client.get(
            f"/api/v1/chat/attachments/{attachments[0]['id']}", headers=_headers()
        )
        assert download.status_code == 200
        assert "attachment" in download.headers["content-disposition"]
        assert download.content == generated.read_bytes()


def test_reply_without_generated_file_has_no_attachments_key(web_settings, tmp_path):
    web_settings.features.memory = False

    def responder(_prompt, *, fn_chunk, **_kwargs):
        fn_chunk("So texto")
        return "So texto"

    hub = EventHub()
    coordinator = ChatCoordinator(
        settings=web_settings,
        event_hub=hub,
        conversation_manager=ConversationManager(tmp_path / "conversations"),
        responder=responder,
        ensure_model_ready=lambda _status: None,
    )
    app = create_app(settings=web_settings, event_hub=hub, chat_coordinator=coordinator)

    with TestClient(app) as client:
        accepted = client.post(
            "/api/v1/chat/messages", headers=_headers(), json={"message": "oi"}
        ).json()
        job = _wait_for_job(client, accepted["job"]["id"])

        conversation = client.get(
            f"/api/v1/chat/conversations/{job['conversation_id']}", headers=_headers()
        ).json()["conversation"]
        reply = [m for m in conversation["messages"] if m["role"] == "assistant"][-1]
        assert "attachments" not in reply["metadata"]
        assert job["attachments"] == []


def test_download_route_requires_a_real_attachment(web_settings, tmp_path):
    web_settings.features.memory = False
    hub = EventHub()
    coordinator = ChatCoordinator(
        settings=web_settings,
        event_hub=hub,
        conversation_manager=ConversationManager(tmp_path / "conversations"),
        responder=lambda *_a, **_k: "oi",
        ensure_model_ready=lambda _status: None,
    )
    app = create_app(settings=web_settings, event_hub=hub, chat_coordinator=coordinator)

    with TestClient(app) as client:
        assert (
            client.get("/api/v1/chat/attachments/" + "a" * 32, headers=_headers()).status_code
            == 404
        )
        assert (
            client.get("/api/v1/chat/attachments/nao-existe", headers=_headers()).status_code == 404
        )


def test_download_route_is_protected(web_settings, tmp_path):
    web_settings.features.memory = False
    hub = EventHub()
    coordinator = ChatCoordinator(
        settings=web_settings,
        event_hub=hub,
        conversation_manager=ConversationManager(tmp_path / "conversations"),
        responder=lambda *_a, **_k: "oi",
        ensure_model_ready=lambda _status: None,
    )
    app = create_app(settings=web_settings, event_hub=hub, chat_coordinator=coordinator)

    with TestClient(app) as client:
        assert client.get("/api/v1/chat/attachments/" + "a" * 32).status_code == 401


# --- document tool integration -------------------------------------------


def test_document_tool_announces_the_attached_file(web_settings, tmp_path, monkeypatch):
    from ai import tools

    form = _docx(tmp_path / "modelo.docx")
    document = Document(str(form))
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nome da criança:"
    table.cell(0, 1).text = ""
    document.save(form)

    store = _store(tmp_path)
    monkeypatch.setattr(tools, "_allowed_file_roots", lambda: [tmp_path])
    monkeypatch.setattr(tools, "_resolve", lambda value: Path(value), raising=False)
    monkeypatch.setattr(tools, "get_settings", lambda: web_settings)

    with collect_outputs(store) as produced:
        text = tools._tool_preencher_documento(str(form), {"Nome da criança": "ARTHUR"})

    assert produced
    assert "anexado para download" in text
    assert store.get(produced[0]["id"]).path.is_file()
    working = Path(web_settings.data_dir) / "cache" / "document_outputs"
    assert not list(working.glob("*.docx"))


def test_template_request_delivers_download_even_when_model_only_drafts_in_chat(
    web_settings, tmp_path, monkeypatch
):
    from io import BytesIO
    from types import SimpleNamespace

    import ai.rag as rag
    from ai import react, tools

    web_settings.features.memory = False
    web_settings.decision.enabled = False
    source = _docx(tmp_path / "origem.docx", "Nome: João da Silva")
    template = _docx(tmp_path / "modelo.docx", "{{nome}}")
    fabricated = "Preenchimento sugerido: nome inventado e relatorio apenas na conversa."

    class Model:
        def create_chat_completion(self, **kwargs):
            yield {"choices": [{"delta": {"content": fabricated}}]}

    manager = SimpleNamespace(
        route_and_invoke=lambda *a, **k: ("stub-model", Model()),
        get_last_decision=lambda: None,
        get_current_complexity=lambda: "simples",
    )
    monkeypatch.setattr(react, "get_multi_model_manager", lambda: manager)
    monkeypatch.setattr(react, "get_settings", lambda: web_settings)
    monkeypatch.setattr(tools, "get_settings", lambda: web_settings)
    monkeypatch.setattr(tools, "_allowed_file_roots", lambda: [tmp_path])
    monkeypatch.setattr(react, "_agenda_prompt_context", lambda: "")
    monkeypatch.setattr(rag, "buscar_contexto", lambda *a, **k: [])
    # Approval policy has its own tests; this executor exercises the authorized
    # physical fill and delivery after the model forgot to call the tool.
    monkeypatch.setattr(
        react,
        "executar_ferramenta",
        lambda name, arguments, **kwargs: tools._tool_preencher_documento_com_fontes(
            **arguments, usar_modelo=False
        ),
    )
    hub = EventHub()
    coordinator = ChatCoordinator(
        settings=web_settings,
        event_hub=hub,
        conversation_manager=ConversationManager(tmp_path / "conversations"),
        responder=lambda prompt, **kwargs: react.loop_react(prompt, **kwargs)[0],
        ensure_model_ready=lambda status: None,
    )
    app = create_app(settings=web_settings, event_hub=hub, chat_coordinator=coordinator)
    ids = [
        coordinator.attachments.save(path.name, path.read_bytes()).id for path in [source, template]
    ]
    with TestClient(app) as client:
        accepted = client.post(
            "/api/v1/chat/messages",
            headers=_headers(),
            json={
                "message": "Preencha modelo.docx com as informações de origem.docx",
                "attachment_ids": ids,
                "agent_mode": "documentos",
            },
        ).json()
        job = _wait_for_job(client, accepted["job"]["id"])
        assert job["status"] == "completed", job
        assert fabricated not in job["response"]
        assert job["chunk_count"] == 0
        assert len(job["attachments"]) == 1
        attachment = job["attachments"][0]
        assert attachment["name"] == "modelo - preenchido.docx"
        downloaded = client.get(f"/api/v1/chat/attachments/{attachment['id']}", headers=_headers())
        assert downloaded.status_code == 200
        assert Document(BytesIO(downloaded.content)).paragraphs[0].text == "João da Silva"


def test_delivery_failure_is_reported_instead_of_claiming_a_download(
    web_settings, tmp_path, monkeypatch
):
    from ai import tools

    template = _docx(tmp_path / "modelo.docx", "{{nome}}")
    store = OutputAttachmentStore(
        root=tmp_path / "outputs", allowed_extensions=ALLOWED, max_bytes=1
    )
    monkeypatch.setattr(tools, "get_settings", lambda: web_settings)
    monkeypatch.setattr(tools, "_allowed_file_roots", lambda: [tmp_path])
    with (
        collect_outputs(store) as produced,
        pytest.raises(ValueError, match="nao foi possivel disponibiliza-lo para download"),
    ):
        tools._tool_preencher_documento(str(template), {"nome": "João"})
    assert produced == []
