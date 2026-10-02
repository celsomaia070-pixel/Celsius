import importlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

import ai.react
import ai.tools
from ai.react import (
    _chart_arguments_from_markdown_table,
    _document_was_written,
    _execute_textual_chart_call,
    _extract_textual_tool_call,
    _extract_xml_tool_calls,
    _is_document_fill_request,
    _required_local_tools,
    _response_claims_document_ready,
    _strip_visible_tool_calls,
    _weather_evidence_for_task,
    _inventory_chart_arguments,
    _needs_pedagogical_report_artifact,
    _response_has_generated_chart,
    _is_chart_request,
    _try_direct_business_chart,
    _try_direct_stock_movement,
    _filtrar_ferramentas,
    _try_direct_business_report,
    loop_react,
)
from ai.tools import (
    REGISTRO_FERRAMENTAS,
    _normalize_chart_arguments,
    _tool_criar_editar_arquivo,
    _tool_gerar_documento_local,
    _tool_inspecionar_formulario_documento,
    _tool_preencher_documento,
    _tool_ler_arquivo,
    _validate_path,
    obter_schemas_openai,
)


def test_extracts_xml_tool_call_emitted_by_gguf_template():
    response = """
    <tool_call>
    <function=pesquisar_noticias>
    <parameter=query>
    inteligencia artificial noticias semana passada
    </parameter>
    </function>
    </tool_call>
    """

    assert _extract_xml_tool_calls(response) == [
        ("pesquisar_noticias", {"query": "inteligencia artificial noticias semana passada"})
    ]
    assert _strip_visible_tool_calls(response + "\nResposta visivel") == "Resposta visivel"


def test_weather_task_uses_the_returned_source_instead_of_model_prose():
    source = "[FONTE_WEB] Open-Meteo — previsao para Marilia, São Paulo\nURL: https://api.open-meteo.com/v1/forecast?x=1"
    task = {"steps": [{"tool": "pesquisar_web", "status": "succeeded", "result": source}]}

    assert _weather_evidence_for_task(task, "Previsao do tempo para Marilia") == source
    assert not _weather_evidence_for_task(task, "Pesquise noticias de tecnologia")


def test_local_data_requests_require_authoritative_tools():
    required = _required_local_tools("acesse meu estoque, gere um relatorio e crie um grafico")
    assert required["listar_estoque"] == {}
    assert required["gerar_relatorio_local"]["fonte"] == "Estoque"
    assert "buscar_memoria" not in _required_local_tools("salve esta informacao na memoria")


from core.settings import SecuritySettings, Settings
from core.tool_approval import SENSITIVE_TOOLS


class TestToolRegistry:
    def test_registry_not_empty(self):
        assert len(REGISTRO_FERRAMENTAS) > 0

    def test_required_tools_exist(self):
        names = {f.nome for f in REGISTRO_FERRAMENTAS}
        for required in ["processar_arquivo", "pesquisar_web", "executar_codigo"]:
            assert required in names, f"Required tool '{required}' not found"

    def test_agenda_tools_exist(self):
        names = {f.nome for f in REGISTRO_FERRAMENTAS}

        assert "listar_agenda" in names
        assert "criar_compromisso_agenda" in names
        assert "marcar_lembrete_agenda" in names

    def test_each_tool_has_schema(self):
        for f in REGISTRO_FERRAMENTAS:
            assert f.schema is not None
            assert "properties" in f.schema

    def test_obter_schemas_openai_format(self):
        schemas = obter_schemas_openai()
        assert len(schemas) == len(REGISTRO_FERRAMENTAS)
        for s in schemas:
            assert "function" in s
            assert "name" in s["function"]
            assert "parameters" in s["function"]


class TestChartToolCompatibility:
    def test_normalizes_record_list_emitted_by_local_model(self):
        arguments = {
            "data": (
                '[{"nome": "Alicate", "quantidade": 52, "categoria": "Geral"}, '
                '{"nome": "Chave", "quantidade": 3, "categoria": "Geral"}]'
            ),
            "tipo": "barras",
            "titulo": "Quantidade em estoque",
        }

        normalized = _normalize_chart_arguments(arguments)

        assert normalized["tipo"] == "bar"
        assert normalized["labels"] == '["Alicate", "Chave"]'
        assert normalized["valores"] == "[52, 3]"
        assert "data" not in normalized

    def test_recovers_and_executes_visible_chart_tool_json(self, monkeypatch):
        response = (
            '{"name": "gerar_grafico", "arguments": {'
            '"data": "[{\\"nome\\": \\"Alicate\\", \\"quantidade\\": 52}]", '
            '"tipo": "barras", "titulo": "Estoque"}}'
        )
        captured = {}

        def fake_execute(name, arguments):
            captured["name"] = name
            captured["arguments"] = arguments
            return (
                "Grafico 'bar' gerado com sucesso.\n"
                "Arquivo: C:\\cache\\estoque.png\n"
                "Exiba-o com: ![Grafico - Estoque](C:\\cache\\estoque.png)"
            )

        monkeypatch.setattr(ai.react, "executar_ferramenta", fake_execute)
        monkeypatch.setattr(
            ai.react,
            "_inventory_report_summary",
            lambda: "**Dados confirmados no inventory.json**",
        )
        monkeypatch.setattr(
            ai.react,
            "_chart_response_from_tool_result",
            lambda _result, _title: "![Grafico - Estoque](C:\\cache\\estoque.png)",
        )

        assert _extract_textual_tool_call(response)[0] == "gerar_grafico"
        result = _execute_textual_chart_call(response)

        assert captured["name"] == "gerar_grafico"
        assert captured["arguments"]["labels"] == '["Alicate"]'
        assert captured["arguments"]["valores"] == "[52]"
        assert "![Grafico - Estoque]" in result
        assert "gerar_grafico" not in result

    def test_remote_image_is_not_accepted_as_generated_chart(self):
        response = "![Grafico](https://example.com/estoque.png)"

        assert _response_has_generated_chart(response) is False

    def test_extracts_chart_data_from_markdown_table(self):
        response = """
| Produto | Quantidade | Categoria |
|---|---:|---|
| Alicate | 52 un. | Geral |
| Chave | 3 un. | Geral |
"""

        arguments = _chart_arguments_from_markdown_table(
            "Crie um grafico de barras da quantidade",
            response,
        )

        assert arguments["labels"] == ["Alicate", "Chave"]
        assert arguments["valores"] == [52.0, 3.0]
        assert arguments["tipo"] == "bar"

    def test_builds_inventory_efficiency_indicator(self):
        items = [
            type("Item", (), {"precisa_repor": False})(),
            type("Item", (), {"precisa_repor": False})(),
            type("Item", (), {"precisa_repor": True})(),
            type("Item", (), {"precisa_repor": True})(),
        ]

        arguments, summary = _inventory_chart_arguments(
            "Crie um indicador de eficiencia do estoque",
            items,
        )

        assert arguments["tipo"] == "kpi"
        assert arguments["valores"] == [50.0]
        assert arguments["meta"] == 90
        assert arguments["unidade"] == "%"
        assert "2 de 4" in summary

    def test_inventory_chart_is_generated_without_waiting_for_llm(self, monkeypatch):
        items = [
            type(
                "Item",
                (),
                {
                    "nome": "Alicate",
                    "categoria": "Ferramentas",
                    "quantidade": 52,
                    "estoque_min": 5,
                    "estoque_max": 70,
                    "precisa_repor": False,
                },
            )(),
            type(
                "Item",
                (),
                {
                    "nome": "Chave",
                    "categoria": "Ferramentas",
                    "quantidade": 3,
                    "estoque_min": 10,
                    "estoque_max": 40,
                    "precisa_repor": True,
                },
            )(),
        ]
        service = type("Inventory", (), {"get_all_items": lambda self: items})()
        captured = {}

        inventory_module = importlib.import_module("core.inventory")
        monkeypatch.setattr(inventory_module, "get_inventory_service", lambda: service)

        def fake_execute(name, arguments):
            captured["name"] = name
            captured["arguments"] = arguments
            return "Arquivo: C:\\cache\\chart.png"

        monkeypatch.setattr(ai.react, "executar_ferramenta", fake_execute)
        monkeypatch.setattr(
            ai.react,
            "_chart_response_from_tool_result",
            lambda _result, _title: "![Grafico](C:\\cache\\chart.png)",
        )

        result = _try_direct_business_chart(
            "Crie um grafico de barras mostrando todos os itens do estoque"
        )

        assert captured["name"] == "gerar_grafico"
        assert captured["arguments"]["tipo"] == "bar"
        assert captured["arguments"]["labels"] == '["Alicate", "Chave"]'
        assert "![Grafico]" in result


class TestBusinessDataToolAccess:
    def test_stock_item_named_rosca_is_not_mistaken_for_chart(self):
        assert not _is_chart_request("de baixa em 1 item tranca rosca")
        assert _is_chart_request("gere um grafico de rosca do estoque")

    def test_natural_stock_output_resolves_item_before_llm(self, monkeypatch):
        item = type("StockItem", (), {"id": "tr1", "nome": "Tranca rosca"})()
        inventory_service = type("InventoryService", (), {"get_all_items": lambda _self: [item]})()
        captured = {}

        def fake_execute(name, arguments, **kwargs):
            captured.update(name=name, arguments=arguments, kwargs=kwargs)
            return "Aprovacao solicitada"

        monkeypatch.setattr("core.inventory.get_inventory_service", lambda: inventory_service)
        monkeypatch.setattr(ai.react, "executar_ferramenta", fake_execute)

        result = _try_direct_stock_movement(
            "de saida em 1 item tranca rosca do estoque",
            approval_scope="conversation-1",
        )

        assert result == "Aprovacao solicitada"
        assert captured["name"] == "saida_estoque"
        assert captured["arguments"] == {"item_id": "tr1", "quantidade": 1}
        assert captured["kwargs"] == {
            "require_approval": True,
            "approval_scope": "conversation-1",
        }

    def test_unrelated_question_does_not_expose_business_database_tools(self):
        names = {tool.nome for tool in _filtrar_ferramentas("Explique energia solar")}

        assert not names

    def test_report_request_exposes_inventory_and_report_tools(self):
        names = {tool.nome for tool in _filtrar_ferramentas("Gere um relatorio do estoque em PDF")}

        assert "listar_estoque" in names
        assert "gerar_relatorio_local" in names

    def test_read_tools_cover_local_business_databases(self):
        names = {tool.nome for tool in _filtrar_ferramentas("Mostre meus dados locais")}

        assert {
            "listar_agenda",
            "listar_clientes",
            "listar_estoque",
            "listar_fornecedores",
            "listar_orcamentos",
            "listar_processos_prazos",
            "listar_produtos_servicos",
        } <= names

    def test_stock_report_is_generated_without_model_tool_choice(self, monkeypatch):
        captured = {}

        def fake_execute(name, arguments):
            captured["name"] = name
            captured["arguments"] = arguments
            return "Relatorio gerado localmente."

        monkeypatch.setattr(ai.react, "executar_ferramenta", fake_execute)
        monkeypatch.setattr(ai.react, "_get_report_content_display", lambda _s, _t: "")

        result = _try_direct_business_report("Crie um relatorio do estoque em PDF")

        assert "Relatorio gerado localmente." in result
        assert "Dados confirmados no inventory.json" in result
        assert captured["name"] == "gerar_relatorio_local"
        assert captured["arguments"]["fonte"] == "Estoque"
        assert captured["arguments"]["formato"] == "pdf"

    def test_stock_context_does_not_disable_deterministic_report(self, monkeypatch):
        monkeypatch.setattr(
            ai.react,
            "_try_direct_business_report",
            lambda _question: "Relatorio confirmado pelo inventory.json",
        )
        monkeypatch.setattr(
            ai.react,
            "get_multi_model_manager",
            lambda: pytest.fail("O LLM nao deve ser usado para este relatorio"),
        )

        response, _steps = loop_react(
            {
                "pergunta": "gere um relatorio de meu estoque",
                "documento": "Dados do estoque do usuario (2 itens)",
                "nome_documento": "Dados do Estoque",
            }
        )

        assert response == "Relatorio confirmado pelo inventory.json"


class TestToolPathSecurity:
    def test_validate_relative_path_inside_base_dir(self, tmp_path, monkeypatch):
        allowed_file = tmp_path / "allowed.txt"
        allowed_file.write_text("ok", encoding="utf-8")
        settings = Settings(
            base_dir=tmp_path,
            security=SecuritySettings(allowed_file_roots=(str(tmp_path),)),
        )
        monkeypatch.setattr(ai.tools, "get_settings", lambda: settings)

        assert _validate_path("allowed.txt") == allowed_file.resolve()

    def test_validate_path_denies_outside_allowed_roots(self, tmp_path, monkeypatch):
        allowed = tmp_path / "allowed"
        outside = tmp_path / "outside"
        allowed.mkdir()
        outside.mkdir()
        outside_file = outside / "secret.txt"
        outside_file.write_text("secret", encoding="utf-8")
        settings = Settings(
            base_dir=allowed,
            security=SecuritySettings(allowed_file_roots=(str(allowed),)),
        )
        monkeypatch.setattr(ai.tools, "get_settings", lambda: settings)

        with pytest.raises(PermissionError):
            _validate_path(str(outside_file))


class TestLerArquivoPaginacao:
    def _settings(self, tmp_path, monkeypatch):
        settings = Settings(
            base_dir=tmp_path,
            security=SecuritySettings(allowed_file_roots=(str(tmp_path),)),
        )
        monkeypatch.setattr(ai.tools, "get_settings", lambda: settings)
        return settings

    def test_le_janela_no_meio_do_arquivo(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        arquivo = tmp_path / "notas.txt"
        arquivo.write_text("".join(f"linha {i}\n" for i in range(1, 11)), encoding="utf-8")

        resultado = _tool_ler_arquivo(str(arquivo), inicio=4, linhas=3)

        assert "(linhas 4-6 de 10)" in resultado
        assert "linha 4" in resultado
        assert "linha 6" in resultado
        assert "linha 3" not in resultado

    def test_inicio_alem_do_fim_informa_total(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        arquivo = tmp_path / "notas.txt"
        arquivo.write_text("linha 1\nlinha 2\n", encoding="utf-8")

        resultado = _tool_ler_arquivo(str(arquivo), inicio=50, linhas=10)

        assert "Total de 2 linhas" in resultado

    def test_decodifica_cp1252(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        arquivo = tmp_path / "gravar.txt"
        arquivo.write_bytes("Olá, mundo com ç e ão".encode("cp1252"))

        assert "Olá, mundo com ç e ão" in _tool_ler_arquivo(str(arquivo))

    def test_decodifica_utf16(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        arquivo = tmp_path / "legado.txt"
        arquivo.write_bytes("Conteúdo legado em UTF-16".encode("utf-16"))

        assert "Conteúdo legado em UTF-16" in _tool_ler_arquivo(str(arquivo))

    def test_arquivo_vazio(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        arquivo = tmp_path / "vazio.txt"
        arquivo.write_text("", encoding="utf-8")

        assert "arquivo vazio" in _tool_ler_arquivo(str(arquivo))


class TestCriarEditarArquivo:
    def _settings(self, tmp_path, monkeypatch):
        settings = Settings(
            base_dir=tmp_path,
            security=SecuritySettings(allowed_file_roots=(str(tmp_path),)),
        )
        monkeypatch.setattr(ai.tools, "get_settings", lambda: settings)
        return settings

    def test_registrada_e_sensivel(self):
        names = {f.nome for f in REGISTRO_FERRAMENTAS}
        assert "criar_editar_arquivo" in names
        assert "criar_editar_arquivo" in SENSITIVE_TOOLS

    def test_cria_arquivo(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        destino = tmp_path / "notas.md"

        resultado = _tool_criar_editar_arquivo(str(destino), "Conteudo novo")

        assert "criado" in resultado
        assert destino.read_text(encoding="utf-8") == "Conteudo novo"

    def test_criar_nao_sobrescreve_existente(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        destino = tmp_path / "notas.md"
        destino.write_text("original", encoding="utf-8")

        resultado = _tool_criar_editar_arquivo(str(destino), "novo", modo="criar")

        assert "ja existe" in resultado
        assert destino.read_text(encoding="utf-8") == "original"

    def test_anexar_ao_final(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        destino = tmp_path / "log.txt"
        destino.write_text("primeira\n", encoding="utf-8")

        _tool_criar_editar_arquivo(str(destino), "segunda\n", modo="anexar")

        assert destino.read_text(encoding="utf-8") == "primeira\nsegunda\n"

    def test_sobrescrever_cria_backup(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        destino = tmp_path / "dados.json"
        destino.write_text('{"v": 1}', encoding="utf-8")

        resultado = _tool_criar_editar_arquivo(str(destino), '{"v": 2}', modo="sobrescrever")

        assert "sobrescrito" in resultado
        assert destino.read_text(encoding="utf-8") == '{"v": 2}'
        backup = tmp_path / "dados.json.bak"
        assert backup.read_text(encoding="utf-8") == '{"v": 1}'

    def test_extensao_fora_da_lista_e_negada(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        destino = tmp_path / "script.py"

        resultado = _tool_criar_editar_arquivo(str(destino), "print('oi')")

        assert "Formato nao permitido" in resultado
        assert not destino.exists()

    def test_fora_das_pastas_autorizadas_e_negado(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        fora = tmp_path.parent / "fora.txt"

        resultado = _tool_criar_editar_arquivo(str(fora), "conteudo")

        assert "Acesso negado" in resultado
        assert not fora.exists()

    def test_modo_invalido(self, tmp_path, monkeypatch):
        self._settings(tmp_path, monkeypatch)
        destino = tmp_path / "notas.txt"

        resultado = _tool_criar_editar_arquivo(str(destino), "x", modo="apagar")

        assert "Modo invalido" in resultado
        assert not destino.exists()


class TestPreencherDocumento:
    def _settings(self, tmp_path, monkeypatch):
        settings = Settings(
            base_dir=tmp_path,
            security=SecuritySettings(allowed_file_roots=(str(tmp_path),)),
        )
        monkeypatch.setattr(ai.tools, "get_settings", lambda: settings)

    def test_tools_are_registered_and_fill_requires_confirmation(self):
        names = {tool.nome for tool in REGISTRO_FERRAMENTAS}
        assert {"inspecionar_formulario_documento", "preencher_documento"} <= names
        assert "preencher_documento" in SENSITIVE_TOOLS

    def test_fill_request_exposes_document_tools_even_with_attachment(self):
        names = {
            tool.nome
            for tool in _filtrar_ferramentas(
                "Preencha este formulario com os dados anexados", has_document=True
            )
        }
        assert {
            "inspecionar_formulario_documento",
            "preencher_documento",
            "preencher_documento_com_fontes",
        } <= names

    def test_inspects_and_fills_docx_copy(self, tmp_path, monkeypatch):
        from docx import Document

        self._settings(tmp_path, monkeypatch)
        source = tmp_path / "formulario.docx"
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Nome:"
        table.cell(0, 1).text = "Anterior"
        document.save(source)

        inspection = json.loads(_tool_inspecionar_formulario_documento(str(source)))
        assert inspection["field_count"] == 1
        result = json.loads(_tool_preencher_documento(str(source), {"nome": "Maria"}, ""))
        output = Path(result["output"])
        assert output != source
        assert Document(output).tables[0].cell(0, 1).text == "Maria"


def test_pedagogical_report_uses_document_generator_not_business_report():
    names = {
        tool.nome for tool in _filtrar_ferramentas("Gere um relatorio pedagogico PEI do aluno")
    }

    assert "gerar_documento_local" in names
    assert "gerar_relatorio_local" not in names
    assert _try_direct_business_report("Gere um relatorio PEI para o aluno") is None


def test_generates_real_pdf_document(tmp_path, monkeypatch):
    settings = Settings(
        base_dir=tmp_path,
        security=SecuritySettings(allowed_file_roots=(str(tmp_path),)),
    )
    monkeypatch.setattr(ai.tools, "get_settings", lambda: settings)
    output = tmp_path / "relatorio.pdf"

    result = json.loads(
        _tool_gerar_documento_local(
            "Relatorio pedagogico", "Conteudo confirmado do aluno.", "pdf", str(output)
        )
    )

    assert result["written"] is True
    assert output.is_file() and output.stat().st_size > 0


def test_indexed_pedagogical_report_gets_an_artifact_backstop():
    task = SimpleNamespace(task={"steps": []})
    messages = [
        {
            "role": "user",
            "content": "<documentos_indexados_nao_confiaveis>\n[Fonte: PEI]\nDados\n"
            "</documentos_indexados_nao_confiaveis>",
        }
    ]

    assert _needs_pedagogical_report_artifact("Gere um relatorio PEI para o aluno", messages, task)
    assert not _needs_pedagogical_report_artifact("Gere um relatorio de estoque", messages, task)


class TestDocumentFillHonestyGate:
    """A filled document is a file on disk, not something a model can describe."""

    def test_recognises_a_document_fill_request(self):
        assert _is_document_fill_request("preencha o formulario do Arthur")
        assert _is_document_fill_request("gera o relatorio em PDF da turma")
        assert _is_document_fill_request("completa a ficha de avaliacao")
        # A question that merely mentions a document is not a fill request.
        assert not _is_document_fill_request("qual o nome do arquivo do PEI?")
        assert not _is_document_fill_request("liste os documentos da escola")
        assert not _is_document_fill_request("resuma o documento anexado")

    def test_recognises_a_fabricated_success_claim(self):
        assert _response_claims_document_ready(
            "Pronto! Preenchi o documento e anexei o arquivo para download."
        )
        assert _response_claims_document_ready("PAEE gerado com sucesso.")
        # No claim about a produced artefact -> nothing to correct.
        assert not _response_claims_document_ready("Posso帮他 preencher se quiser.")
        assert not _response_claims_document_ready("O formulario tem 8 campos.")

    def test_requires_a_confirmed_write_not_merely_a_tool_call(self):
        refused = [
            {
                "role": "assistant",
                "tool_calls": [
                    {
                        "id": "c1",
                        "function": {"name": "preencher_documento_com_fontes"},
                    }
                ],
            },
            {"role": "tool", "tool_call_id": "c1", "content": '{"written": false}'},
        ]
        assert not _document_was_written(refused)

        succeeded = [
            {
                "role": "assistant",
                "tool_calls": [
                    {
                        "id": "c2",
                        "function": {"name": "preencher_documento_com_fontes"},
                    }
                ],
            },
            {
                "role": "tool",
                "tool_call_id": "c2",
                "content": '{"written": true, "output": "PAEE.docx"}',
            },
        ]
        assert _document_was_written(succeeded)

    def test_counts_a_generated_report_only_when_its_tool_confirms_the_file(self):
        messages = [
            {
                "role": "assistant",
                "tool_calls": [{"id": "c4", "function": {"name": "gerar_documento_local"}}],
            },
            {
                "role": "tool",
                "tool_call_id": "c4",
                "content": '{"written": true, "output": "relatorio.pdf"}',
            },
        ]

        assert _document_was_written(messages)

    def test_a_tool_call_for_another_tool_does_not_count(self):
        messages = [
            {
                "role": "assistant",
                "tool_calls": [
                    {"id": "c3", "function": {"name": "inspecionar_formulario_documento"}}
                ],
            },
            {
                "role": "tool",
                "tool_call_id": "c3",
                "content": '{"field_count": 8}',
            },
        ]
        assert not _document_was_written(messages)

    def test_loop_blocks_a_hallucinated_document(self, monkeypatch):
        # The exact failure reported by the user: the model narrated success and
        # never called a tool, so no file existed.
        fabricated = (
            "Prontinho! Preenchi o PAEE com os dados do Arthur e anexei o "
            "arquivo preenchido para download."
        )

        class _FakeLlama:
            def create_chat_completion(self, **_kwargs):
                yield {"choices": [{"delta": {"content": fabricated}}]}

        class _FakeManager:
            def route_and_invoke(self, *_args, **_kwargs):
                return "stub-model", _FakeLlama()

            def get_last_decision(self):
                return None

            def get_current_complexity(self):
                return "simples"

        monkeypatch.setattr(ai.react, "get_multi_model_manager", lambda: _FakeManager())
        # The pre-existing local-report gate may still run; what must not happen
        # is a fill tool producing a file the model never asked for.
        calls: list[str] = []
        monkeypatch.setattr(
            ai.react,
            "executar_ferramenta",
            lambda name, *_a, **_k: calls.append(name) or "relatorio local pronto",
        )

        resposta, _steps = loop_react(
            {"pergunta": "preencha o formulario do Arthur e anexe o documento"}
        )

        assert not {"preencher_documento", "preencher_documento_com_fontes"} & set(calls)
        assert fabricated not in resposta
        assert "Nao preenchi nenhum documento" in resposta
        assert "preencher_documento_com_fontes" in resposta

    def test_loop_keeps_an_honest_answer_untouched(self, monkeypatch):
        honest = "Voce poderia me passar o caminho do formulario em branco?"

        class _FakeLlama:
            def create_chat_completion(self, **_kwargs):
                yield {"choices": [{"delta": {"content": honest}}]}

        class _FakeManager:
            def route_and_invoke(self, *_args, **_kwargs):
                return "stub-model", _FakeLlama()

            def get_last_decision(self):
                return None

            def get_current_complexity(self):
                return "simples"

        monkeypatch.setattr(ai.react, "get_multi_model_manager", lambda: _FakeManager())

        resposta, _steps = loop_react(
            {"pergunta": "preencha o formulario do Arthur e anexe o documento"}
        )

        assert honest in resposta
