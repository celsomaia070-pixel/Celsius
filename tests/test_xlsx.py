import os
import tempfile
from datetime import date

from openpyxl import Workbook, load_workbook

from processors import PROCESSADORES, processar_arquivo
from processors.report import GeradorRelatorio
from processors.xlsx import ProcessadorXLSX


def _criar_workbook(caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    ws.append(["Produto", "Qtd", "Valor", "Entrega"])
    ws.append(["Alicate", 52, 129.5, date(2026, 9, 18)])
    ws.append(["Chave", 3, 12.0, None])
    wb.save(caminho)


def test_processador_xlsx_registrado():
    assert PROCESSADORES[".xlsx"] is ProcessadorXLSX


def test_processar_planilha(tmp_path):
    arquivo = tmp_path / "vendas.xlsx"
    _criar_workbook(arquivo)

    texto = processar_arquivo(str(arquivo), base_dir=tmp_path)

    assert "Planilha: Vendas" in texto
    assert "Produto | Qtd | Valor | Entrega" in texto
    assert "Alicate" in texto
    assert "129.5" in texto
    assert "12" in texto
    assert "2026-09-18" in texto


def test_processar_planilha_vazia(tmp_path):
    arquivo = tmp_path / "vazia.xlsx"
    wb = Workbook()
    wb.save(arquivo)

    assert "vazia" in processar_arquivo(str(arquivo), base_dir=tmp_path).lower()


def test_exportar_xlsx_gera_arquivo_valido(tmp_path):
    saida = str(tmp_path / "relatorio.xlsx")
    GeradorRelatorio.exportar_xlsx(
        "Relatorio",
        "# Resumo\n## Vendas\n| Item | Valor |\n| A | 10 |\n| B | 20 |",
        saida,
        {"Fonte": "Estoque"},
    )

    assert os.path.exists(saida)
    wb = load_workbook(saida)
    ws = wb.active
    assert ws["A1"].value == "Relatorio"
    valores = {cell.value for row in ws.iter_rows() for cell in row}
    assert "Resumo" in valores
    assert "Item" in valores
    assert "10" in valores
