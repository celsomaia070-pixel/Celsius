from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf import PdfReader
from pypdf.generic import (
    ArrayObject,
    BooleanObject,
    DictionaryObject,
    NameObject,
    NumberObject,
    RectangleObject,
    TextStringObject,
)

from core.document_forms import fill_document_form, inspect_document_form


def _word_form(path: Path) -> None:
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Nome da criança:"
    table.cell(0, 1).text = "Nome anterior"
    table.cell(1, 0).text = "Data de nascimento:"
    table.cell(1, 1).text = ""
    merged = table.cell(2, 0).merge(table.cell(2, 1))
    merged.text = "Possui alguma necessidade específica? Resposta anterior"
    document.add_paragraph("Professora: {{professora_aee}}")
    document.save(path)


def test_inspects_and_fills_word_table_without_overwriting_original(tmp_path: Path):
    source = tmp_path / "modelo.docx"
    output = tmp_path / "preenchido.docx"
    _word_form(source)

    inspection = inspect_document_form(source)
    labels = {field["label"] for field in inspection["fields"]}
    assert "Nome da criança:" in labels
    assert "Possui alguma necessidade específica?" in labels
    assert "professora_aee" in labels

    result = fill_document_form(
        source,
        {
            "nome_da_crianca": "João Lucas",
            "Data de nascimento": "13/02/2017",
            "possui alguma necessidade especifica": "Necessita mediação visual.",
            "professora_aee": "Érica Maia",
        },
        output,
    )

    assert result["original_preserved"] is True
    assert result["applied_count"] == 4
    assert result["unmatched"] == []
    original = Document(source)
    filled = Document(output)
    assert original.tables[0].cell(0, 1).text == "Nome anterior"
    assert filled.tables[0].cell(0, 1).text == "João Lucas"
    assert filled.tables[0].cell(1, 1).text == "13/02/2017"
    assert "Necessita mediação visual." in filled.tables[0].cell(2, 0).text
    assert "Érica Maia" in filled.paragraphs[-1].text


def test_never_overwrites_source_or_existing_output(tmp_path: Path):
    source = tmp_path / "modelo.docx"
    _word_form(source)
    with pytest.raises(ValueError, match="nunca e sobrescrito"):
        fill_document_form(source, {"nome_da_crianca": "Outro"}, source)

    existing = tmp_path / "saida.docx"
    existing.write_bytes(source.read_bytes())
    with pytest.raises(FileExistsError, match="ja existe"):
        fill_document_form(source, {"nome_da_crianca": "Outro"}, existing)


def _checkbox_form(path: Path, *, split_paragraphs: bool) -> None:
    """Build the two real shapes: label cell + option cell, and one shared cell."""

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Etapa de Ensino:"
    options = table.cell(0, 1)
    if split_paragraphs:
        # Word often puts one option per paragraph, which _clean flattens.
        for index, text in enumerate(
            (
                "(  ) Educação Infantil",
                "( X ) Ensino Fundamental ",
                "(  ) Programas / Projetos Pedagógicos Específicos",
            )
        ):
            paragraph = options.paragraphs[0] if index == 0 else options.add_paragraph()
            paragraph.add_run(text)
    else:
        options.text = "(X) Sala Multifuncacional () Apoio Colaborativo ( ) Outro"
    # Label and options share a single cell, as in the real AEE form.
    shared = table.cell(1, 0).merge(table.cell(1, 1))
    shared.text = "  Composição do atendimento: (  ) individual         (X) compartilhado "
    document.save(path)


def test_checkbox_group_toggles_marker_without_duplicating_text(tmp_path: Path):
    source = tmp_path / "checks.docx"
    output = tmp_path / "checks_preenchido.docx"
    _checkbox_form(source, split_paragraphs=True)

    result = fill_document_form(
        source,
        {"Etapa de Ensino": "Educação Infantil", "Composição do atendimento": "individual"},
        output,
    )

    assert result["unmatched"] == []
    assert result["needs_review"] == []

    filled = Document(output)
    options_cell = filled.tables[0].cell(0, 1)
    text = options_cell.text
    # The requested option is marked and the previous one is cleared.
    assert "(X) Educação Infantil" in text
    assert "( ) Ensino Fundamental" in text
    # The caption must not be appended a second time.
    assert text.count("Educação Infantil") == 1
    assert text.count("Ensino Fundamental") == 1
    # Layout is untouched: still three paragraphs, one per option.
    assert len(options_cell.paragraphs) == 3

    shared = filled.tables[0].cell(1, 0)
    assert "(X) individual" in shared.text
    assert "( ) compartilhado" in shared.text
    assert shared.text.count("individual") == 1


def test_checkbox_options_are_reported_by_inspection(tmp_path: Path):
    source = tmp_path / "checks.docx"
    _checkbox_form(source, split_paragraphs=True)

    inspection = inspect_document_form(source)
    groups = {
        field["label"]: field for field in inspection["fields"] if field["kind"] == "checkbox"
    }

    etapa = groups["Etapa de Ensino"]
    assert etapa["options"] == [
        "Educação Infantil",
        "Ensino Fundamental",
        "Programas / Projetos Pedagógicos Específicos",
    ]
    assert etapa["checked"] == ["Ensino Fundamental"]

    composicao = groups["Composição do atendimento"]
    assert composicao["options"] == ["individual", "compartilhado"]
    assert composicao["checked"] == ["compartilhado"]


def test_unknown_checkbox_option_is_reported_instead_of_silently_ignored(tmp_path: Path):
    source = tmp_path / "checks.docx"
    output = tmp_path / "out.docx"
    _checkbox_form(source, split_paragraphs=True)

    result = fill_document_form(source, {"Etapa de Ensino": "Educação Superior"}, output)

    # It must not claim success, and the caller has to see why.
    assert "Etapa de Ensino" in result["unmatched"]
    assert any("nenhuma opcao" in note for note in result["needs_review"])


def test_run_formatting_is_preserved_when_filling(tmp_path: Path):
    source = tmp_path / "fmt.docx"
    output = tmp_path / "fmt_out.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Nome da criança:"
    run = table.cell(0, 1).paragraphs[0].add_run("Nome anterior")
    run.font.size = 152400
    run.bold = True
    document.save(source)

    fill_document_form(source, {"Nome da criança": "João"}, output)

    filled_run = [
        r for r in Document(output).tables[0].cell(0, 1).paragraphs[0].runs if r.text.strip()
    ][-1]
    assert filled_run.text == "João"
    assert filled_run.font.size == 152400
    assert filled_run.bold is True


def test_static_pdf_is_identified_and_not_written_over(tmp_path: Path):
    source = tmp_path / "estatico.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with source.open("wb") as handle:
        writer.write(handle)

    inspection = inspect_document_form(source)
    assert inspection["form_type"] == "pdf_static"
    assert inspection["fillable"] is False
    with pytest.raises(ValueError, match="nao possui campos editaveis"):
        fill_document_form(source, {"nome": "João"}, tmp_path / "saida.pdf")
    assert not (tmp_path / "saida.pdf").exists()


def test_fillable_pdf_accepts_the_human_label_and_updates_the_acroform(tmp_path: Path):
    source = tmp_path / "formulario.pdf"
    output = tmp_path / "preenchido.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    field = DictionaryObject(
        {
            NameObject("/FT"): NameObject("/Tx"),
            NameObject("/T"): TextStringObject("student_name"),
            NameObject("/TU"): TextStringObject("Nome do aluno"),
            NameObject("/V"): TextStringObject("Anterior"),
            NameObject("/Ff"): NumberObject(0),
            NameObject("/Subtype"): NameObject("/Widget"),
            NameObject("/Type"): NameObject("/Annot"),
            NameObject("/Rect"): RectangleObject([20, 220, 280, 250]),
        }
    )
    field_ref = writer._add_object(field)
    page[NameObject("/Annots")] = ArrayObject([field_ref])
    acroform = DictionaryObject(
        {
            NameObject("/Fields"): ArrayObject([field_ref]),
            NameObject("/NeedAppearances"): BooleanObject(True),
        }
    )
    writer._root_object[NameObject("/AcroForm")] = writer._add_object(acroform)
    with source.open("wb") as handle:
        writer.write(handle)

    inspection = inspect_document_form(source)
    assert inspection["form_type"] == "pdf_acroform"
    assert inspection["fields"][0]["label"] == "Nome do aluno"

    result = fill_document_form(source, {"Nome do aluno": "Novo nome"}, output)

    assert result["unmatched"] == []
    assert PdfReader(output).get_fields()["student_name"]["/V"] == "Novo nome"


def test_two_column_content_table_is_not_mistaken_for_form_fields(tmp_path: Path):
    source = tmp_path / "plano.docx"
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Professor do AEE"
    table.cell(0, 1).text = "Maria"
    table.cell(1, 0).text = "Objetivos"
    table.cell(1, 1).text = "Estrategias/Recursos"
    table.cell(2, 0).text = "Desenvolver a coordenacao motora fina."
    table.cell(2, 1).text = "Atividades de tracado e desenho."
    document.save(source)

    labels = {field["label"] for field in inspect_document_form(source)["fields"]}

    assert "Professor do AEE" in labels
    assert "Objetivos" not in labels
    assert "Desenvolver a coordenacao motora fina." not in labels


def test_blank_merged_section_heading_is_not_a_fillable_field(tmp_path: Path):
    source = tmp_path / "secoes.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Competencias Cognitivas:"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Quais sao os interesses?"
    document.save(source)

    labels = {field["label"] for field in inspect_document_form(source)["fields"]}

    assert "Competencias Cognitivas:" not in labels
    assert "Quais sao os interesses?" in labels
