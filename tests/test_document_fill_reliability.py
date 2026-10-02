"""Regressions for evidence validation and DOCX package fidelity."""

import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.agent_artifacts import verify_task, workspace_for
from core.chat_attachments import retain_task_document_attachments
from core.document_forms import fill_document_form
from core.document_pipeline import fill_from_sources
from core.tool_policy import Risk, assess_tool


def pair(path, label, value):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = label
    table.cell(0, 1).text = value
    document.save(path)
    return path


def test_semantic_reference_copies_only_a_real_source_value(tmp_path, monkeypatch):
    source = pair(tmp_path / "origem.docx", "Anotação", "Participa de atividades com apoio visual.")
    template = pair(tmp_path / "modelo.docx", "Observações:", "")
    monkeypatch.setattr("core.document_pipeline._ask_llm", lambda *args: (
        {"Observações:": {"origem": "anotacao", "confianca": 0.95, "valor": "Inventado"}}, ""
    ))
    result = fill_from_sources(template, [source], tmp_path / "saida.docx")
    assert Document(result["output"]).tables[0].cell(0, 1).text == "Participa de atividades com apoio visual."
    assert result["mapping"][0]["evidence"][0]["source"] == str(source.resolve())
    assert result["mapping"][0]["confidence"] == 0.95


def test_model_cannot_supply_free_values_or_unknown_or_low_confidence_references(tmp_path, monkeypatch):
    source = pair(tmp_path / "origem.docx", "Anotação", "Valor confirmado")
    template = pair(tmp_path / "modelo.docx", "Observações:", "")
    for proposal in (
        {"Observações:": "Dado inventado"},
        {"Observações:": {"origem": "inexistente", "confianca": 1}},
        {"Observações:": {"origem": "anotacao", "confianca": 0.5}},
    ):
        monkeypatch.setattr("core.document_pipeline._ask_llm", lambda *args, p=proposal: (p, ""))
        result = fill_from_sources(template, [source], tmp_path / "saida.docx")
        assert result["written"] is False
        assert result["rejected"]
        assert not (tmp_path / "saida.docx").exists()


def test_cross_document_conflict_does_not_select_first_name(tmp_path):
    a = pair(tmp_path / "a.docx", "Nome do aluno", "João")
    b = pair(tmp_path / "b.docx", "Nome do aluno", "Pedro")
    target = pair(tmp_path / "modelo.docx", "Nome:", "")
    result = fill_from_sources(target, [a, b], use_llm=False)
    assert result["written"] is False
    assert result["conflicts"]["nome do aluno"] == ["João", "Pedro"]
    assert "Nome:" in result["still_open"]


def test_preview_does_not_create_any_output(tmp_path):
    source = pair(tmp_path / "origem.docx", "Nome do aluno", "João da Silva")
    target = pair(tmp_path / "modelo.docx", "Nome:", "")
    before = target.read_bytes()
    result = fill_from_sources(target, [source], tmp_path / "nao_criar.docx", use_llm=False, preview_only=True)
    assert result["values"]["Nome:"] == "João da Silva"
    assert result["written"] is False
    assert not (tmp_path / "nao_criar.docx").exists()
    assert target.read_bytes() == before


def test_package_parts_and_run_objects_survive_targeted_edits(tmp_path):
    source = tmp_path / "modelo.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Cabeçalho preservado"
    document.sections[0].footer.paragraphs[0].text = "Rodapé preservado"
    p = document.add_paragraph()
    run = p.add_run("{{nome}}")
    run.bold = True
    # A non-text run child would be destroyed by assigning Run.text.
    drawing = OxmlElement("w:drawing")
    run._r.append(drawing)
    document.save(source)
    result = fill_document_form(source, {"nome": "João"}, tmp_path / "saida.docx")
    with ZipFile(source) as original, ZipFile(result["output"]) as filled:
        assert original.namelist() == filled.namelist()
        for name in original.namelist():
            if name != "word/document.xml":
                assert original.read(name) == filled.read(name), name
    output = Document(result["output"])
    assert output.paragraphs[0].runs[0].bold
    assert output.paragraphs[0]._p.xpath(".//w:drawing")
    assert output.paragraphs[0].text == "João"


def test_nested_table_and_header_fields_are_filled(tmp_path):
    target = tmp_path / "modelo.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Escola: {{escola}}"
    outer = document.add_table(rows=1, cols=1)
    nested = outer.cell(0, 0).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Nome:"
    document.save(target)
    source = pair(tmp_path / "origem.docx", "Nome do aluno", "João")
    source_document = Document(source)
    source_document.add_paragraph("Escola: Escola Exemplo")
    source_document.save(source)
    result = fill_from_sources(target, [source], tmp_path / "saida.docx", use_llm=False)
    assert Document(result["output"]).tables[0].cell(0, 0).tables[0].cell(0, 1).text == "João"
    assert Document(result["output"]).sections[0].header.paragraphs[0].text == "Escola: Escola Exemplo"


def test_named_word_content_control_is_filled_without_rebuilding(tmp_path):
    target = tmp_path / "modelo.docx"
    document = Document()
    control = OxmlElement("w:sdt")
    props = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "Nome")
    props.append(alias)
    content = OxmlElement("w:sdtContent")
    p, r, t = OxmlElement("w:p"), OxmlElement("w:r"), OxmlElement("w:t")
    t.text = "Escolha um nome"
    r.append(t)
    p.append(r)
    content.append(p)
    control.extend([props, content])
    document.element.body.insert(0, control)
    document.save(target)
    result = fill_document_form(target, {"Nome": "João"}, tmp_path / "saida.docx")
    assert "Nome" in result["applied"]
    assert Document(result["output"]).element.xpath(".//w:sdt//w:t")[0].text == "João"


def test_retained_task_inputs_survive_upload_cleanup_without_counting_as_outputs(tmp_path):
    source = pair(tmp_path / "upload.docx", "Nome", "João")
    workspace = workspace_for(tmp_path, "a1b2c3")
    prompt = {"documentos_anexados": [{"nome": "origem.docx", "caminho": str(source)}]}
    retained = retain_task_document_attachments(prompt, workspace)
    source.unlink()
    assert Path(retained["documentos_anexados"][0]["caminho"]).is_file()
    assert verify_task({"id": "a1b2c3", "steps": []}, tmp_path)["artifact_count"] == 0


def test_filename_race_never_removes_or_overwrites_another_tasks_output(tmp_path, monkeypatch):
    import core.document_forms as forms

    target = pair(tmp_path / "modelo.docx", "Nome:", "")
    output = tmp_path / "saida.docx"
    original_writer = forms._fill_docx

    def competing_writer(source, staged, values):
        result = original_writer(source, staged, values)
        output.write_bytes(b"arquivo de outra tarefa")
        return result

    monkeypatch.setattr(forms, "_fill_docx", competing_writer)
    with pytest.raises(FileExistsError):
        fill_document_form(target, {"Nome": "João"}, output)
    assert output.read_bytes() == b"arquivo de outra tarefa"
    assert not list(tmp_path.glob(".celsius-fill-*"))


def test_checkbox_requires_a_related_source_label_not_just_the_word_yes(tmp_path):
    target = pair(tmp_path / "modelo.docx", "Necessita adaptação:", "( ) Sim ( ) Não")
    unrelated = pair(tmp_path / "origem.docx", "Autorização de passeio", "Sim")
    result = fill_from_sources(target, [unrelated], use_llm=False)
    assert result["written"] is False
    assert "Necessita adaptação" in result["still_open"]


def test_checked_source_option_has_traceable_evidence(tmp_path):
    source = pair(tmp_path / "origem.docx", "Necessita adaptação", "(X) Sim ( ) Não")
    target = pair(tmp_path / "modelo.docx", "Necessita adaptação:", "( ) Sim ( ) Não")
    result = fill_from_sources(target, [source], tmp_path / "saida.docx", use_llm=False)
    assert Document(result["output"]).tables[0].cell(0, 1).text == "(X) Sim ( ) Não"
    assert result["mapping"][0]["evidence"][0]["value"] == "Sim"


@pytest.mark.parametrize("answer", ["{{nome}}", "__________", "( ) Sim ( ) Não"])
def test_empty_source_placeholders_are_not_facts(tmp_path, answer):
    source = pair(tmp_path / "origem.docx", "Nome", answer)
    target = pair(tmp_path / "modelo.docx", "Nome:", "")
    result = fill_from_sources(target, [source], use_llm=False)
    assert result["written"] is False


def test_mapping_preview_policy_requires_an_actual_boolean():
    preview = assess_tool("preencher_documento_com_fontes", {"somente_analisar": True})
    assert preview.risk is Risk.READ
    assert preview.requires_confirmation is False
    assert assess_tool("preencher_documento_com_fontes", {"somente_analisar": "true"}).requires_confirmation
    assert assess_tool("preencher_documento_com_fontes", {"somente_analisar": False}).requires_confirmation


def test_artifact_verifier_counts_only_written_document_outputs(tmp_path):
    source = pair(tmp_path / "origem.docx", "Nome", "João")
    output = pair(tmp_path / "saida.docx", "Nome", "João")
    task = {"id": "a1b2c3", "steps": [{
        "tool": "preencher_documento_com_fontes",
        "result": json.dumps({"source": str(source), "output": "", "written": False}),
    }]}
    assert verify_task(task, tmp_path / "data")["artifact_count"] == 0
    task["steps"][0]["result"] = json.dumps({
        "source": str(source), "output": str(output), "written": True,
    }) + "\nDocumento anexado para download: saida.docx"
    report = verify_task(task, tmp_path / "data")
    assert report["artifact_count"] == 1
    assert report["artifacts"][0]["path"] == str(output)


def test_repeated_semantic_destination_does_not_choose_the_last_reference(tmp_path, monkeypatch):
    from core.document_pipeline import _ask_llm
    import core.model_router as router

    payload = {"mapeamentos": [
        {"destino": "Observações:", "origem": "anotacao", "confianca": 0.95},
        {"destino": "Observacoes", "origem": "outro", "confianca": 0.95},
    ]}
    manager = SimpleNamespace(is_healthy=lambda: True, create_chat_completion=lambda **kwargs: {
        "choices": [{"message": {"content": json.dumps(payload)}}],
    })
    monkeypatch.setattr(router, "get_multi_model_manager", lambda: SimpleNamespace(main_manager=manager))
    proposal, _ = _ask_llm([], {"anotacao": "Um", "outro": "Outro"})
    assert proposal["Observações:"]["origem"] is None


def test_same_row_inline_labels_do_not_overwrite_the_next_field(tmp_path):
    source = tmp_path / "origem.docx"
    template = tmp_path / "modelo.docx"
    for path, texts in [(source, ["Nome:\nJoão", "Escola:\nEscola Exemplo"]),
                        (template, ["Nome:", "Escola:"])]:
        document = Document()
        table = document.add_table(rows=1, cols=2)
        for index, text in enumerate(texts):
            table.cell(0, index).text = text
        document.save(path)
    result = fill_from_sources(template, [source], tmp_path / "saida.docx", use_llm=False)
    cells = Document(result["output"]).tables[0].rows[0].cells
    assert cells[0].text == "Nome: João"
    assert cells[1].text == "Escola: Escola Exemplo"


def test_numbered_sections_and_following_answer_rows_are_filled_from_real_source(tmp_path):
    source = tmp_path / "origem.docx"
    template = tmp_path / "modelo.docx"
    for path, populated in [(source, True), (template, False)]:
        document = Document()
        objectives = document.add_table(rows=3, cols=1)
        objectives.cell(0, 0).text = "OBJETIVOS:"
        objectives.cell(1, 0).text = "1- Reconhecer letras." if populated else "1"
        objectives.cell(2, 0).text = "2- Contar objetos." if populated else "2"
        period = document.add_table(rows=2, cols=1)
        period.cell(0, 0).text = "TEMPO ESTIMADO:"
        period.cell(1, 0).text = "Dois meses de 2025: outubro e novembro." if populated else ""
        report = document.add_table(rows=1, cols=1)
        report.cell(0, 0).text = "RELATÓRIO: Participa com apoio visual." if populated else "RELATÓRIO:"
        document.save(path)
    result = fill_from_sources(template, [source], tmp_path / "saida.docx", use_llm=False)
    output = Document(result["output"])
    assert output.tables[0].cell(0, 0).text == "OBJETIVOS:"
    assert output.tables[0].cell(1, 0).text == "1- Reconhecer letras."
    assert output.tables[0].cell(2, 0).text == "2- Contar objetos."
    assert output.tables[1].cell(1, 0).text == "Dois meses de 2025: outubro e novembro."
    assert "Participa com apoio visual." in output.tables[2].cell(0, 0).text
    assert result["still_open"] == []


def test_short_existing_answer_never_replaces_letters_inside_the_label(tmp_path):
    source = tmp_path / "origem.docx"
    template = tmp_path / "modelo.docx"
    for path, answer in [(source, "Confirmado"), (template, "T")]:
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = f"TIPO DE ATENDIMENTO:\n{answer}"
        document.save(path)
    result = fill_from_sources(template, [source], tmp_path / "saida.docx", use_llm=False)
    assert Document(result["output"]).tables[0].cell(0, 0).text == "TIPO DE ATENDIMENTO:\nConfirmado"


def test_semantic_preview_does_not_block_the_required_physical_write():
    from ai.react import _template_write_attempted

    messages = [{"role": "assistant", "tool_calls": [{"function": {
        "name": "preencher_documento_com_fontes", "arguments": '{"somente_analisar": true}',
    }}]}]
    assert _template_write_attempted(messages) is False
    messages[0]["tool_calls"][0]["function"]["arguments"] = '{"somente_analisar": false}'
    assert _template_write_attempted(messages) is True


def test_explicit_destination_filename_resolves_the_template_without_guessing(tmp_path):
    from ai.react import _required_template_arguments

    template = pair(tmp_path / "ficha.docx", "Nome:", "")
    source = pair(tmp_path / "dados.docx", "Nome:", "João")
    arguments = _required_template_arguments({
        "pergunta": "Preencha ficha.docx com as informações de dados.docx",
        "documentos_anexados": [{"nome": path.name, "caminho": str(path)} for path in [template, source]],
    })
    assert arguments["caminho_modelo"] == str(template)
    assert arguments["caminhos_fontes"] == [str(source)]
